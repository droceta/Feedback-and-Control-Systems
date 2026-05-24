import math
import os
from datetime import datetime

import simpleeval
import wikipedia
from docx import Document
from langchain_core.tools import tool
from langchain_openai import ChatOpenAI
from langgraph.prebuilt import create_react_agent


# ================================================================
# 1. CONNECT TO THE LOCAL LM STUDIO SERVER
# ================================================================
local_chat_model = ChatOpenAI(
    base_url="http://192.168.56.1:1234/v1",
    api_key="lm-studio",
    model="google/gemma-4-e2b",
    temperature=0.25,
)


# ================================================================
# 2. DOCUMENT SAVING UTILITY
# ================================================================
def save_report_as_docx(text: str, filename: str) -> str:
    """Save generated text into a formatted Word document."""
    output_folder = "agent_outputs"
    os.makedirs(output_folder, exist_ok=True)

    file_path = os.path.join(output_folder, filename)

    document = Document()
    document.add_heading("Local AI Agent Report", level=0)
    document.add_paragraph(f"Created: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}")

    for paragraph in text.splitlines():
        cleaned = paragraph.strip()
        if cleaned:
            document.add_paragraph(cleaned)

    document.save(file_path)
    return f"Done. I saved the Word document here: {file_path}"


# ================================================================
# 3. TOOLS AVAILABLE TO THE AGENT
# ================================================================
@tool
def math_solver(expression: str) -> str:
    """Evaluate a math expression, including sqrt, pow, abs, pi, and e."""
    try:
        evaluator = simpleeval.SimpleEval()
        evaluator.functions = {
            "sqrt": math.sqrt,
            "pow": math.pow,
            "abs": abs,
        }
        evaluator.names = {
            "pi": math.pi,
            "e": math.e,
        }

        answer = evaluator.eval(expression)
        return f"The answer is {answer}"
    except simpleeval.InvalidExpression as error:
        return f"That math expression is not valid: {error}"
    except Exception as error:
        return f"I could not solve that expression. Reason: {error}"


@tool
def encyclopedia_summary(topic: str) -> str:
    """Get a short Wikipedia-based explanation of a general topic."""
    wikipedia.set_lang("en")

    try:
        return wikipedia.summary(topic, sentences=4, auto_suggest=True)
    except wikipedia.exceptions.DisambiguationError as error:
        choices = ", ".join(error.options[:5])
        return f"The topic '{topic}' has multiple matches. Try one of these: {choices}."
    except wikipedia.exceptions.PageError:
        return f"I could not find a Wikipedia page for '{topic}'."
    except Exception as error:
        return f"Wikipedia search failed for '{topic}'. Reason: {error}"


@tool
def create_custom_report(topic_or_text: str) -> str:
    """Create a Word document from a topic or text supplied by the user."""
    writing_prompt = (
        "Create a short school-style report about the following request. "
        "Use a title, short sections, and simple bullet-style points where useful.\n\n"
        f"Request: {topic_or_text}"
    )
    generated_text = local_chat_model.invoke(writing_prompt).content

    file_stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return save_report_as_docx(generated_text, f"agent_report_{file_stamp}.docx")


@tool
def wikipedia_report_builder(topic: str) -> str:
    """Build a Word document using a longer Wikipedia summary as source material."""
    wikipedia.set_lang("en")

    try:
        summary = wikipedia.summary(topic, sentences=10, auto_suggest=True)
    except wikipedia.exceptions.DisambiguationError as error:
        choices = ", ".join(error.options[:5])
        return f"I cannot make the report yet because '{topic}' is too broad. Try: {choices}."
    except wikipedia.exceptions.PageError:
        return f"I could not find a Wikipedia page for '{topic}'."
    except Exception as error:
        return f"I could not collect Wikipedia information for '{topic}'. Reason: {error}"

    facts = [sentence.strip() for sentence in summary.split(". ") if sentence.strip()]
    opening = ". ".join(facts[:3])
    remaining_facts = "\n".join(f"- {fact}." for fact in facts[3:])

    report_text = f"""
RESEARCH REPORT: {topic.title()}
Generated on {datetime.now().strftime('%B %d, %Y')}

Overview
{opening}.

Important Details
{remaining_facts}

Source Notes
- Main source: Wikipedia summary
- Created using a local agent connected to LM Studio
"""

    file_stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return save_report_as_docx(report_text, f"wiki_research_{file_stamp}.docx")


agent_tools = [
    math_solver,
    encyclopedia_summary,
    create_custom_report,
    wikipedia_report_builder,
]


# ================================================================
# 4. AGENT INSTRUCTIONS
# ================================================================
agent_prompt = """
You are a local AI assistant running through VS Code and LM Studio.

You can choose from these tools:
- Use math_solver for arithmetic and math expressions.
- Use encyclopedia_summary when the user asks for a quick explanation of a topic.
- Use create_custom_report when the user asks for a new Word document from a topic or raw text.
- Use wikipedia_report_builder when the user asks to turn a Wikipedia topic into a Word document.

After using a tool, clearly tell the user what you did. If a document was created,
mention the saved file path.
"""

agent = create_react_agent(
    model=local_chat_model,
    tools=agent_tools,
    prompt=agent_prompt,
)


# ================================================================
# 5. TERMINAL CHAT PROGRAM
# ================================================================
if __name__ == "__main__":
    print("=" * 54)
    print("Local Agentic AI Workspace")
    print("Connected to LM Studio through a local network address.")
    print("Type 'exit' when you are done.")
    print("=" * 54)

    while True:
        user_message = input("\nYou: ").strip()

        if user_message.lower() in {"exit", "quit"}:
            print("Session ended. Goodbye!")
            break

        if not user_message:
            continue

        print("\nWorking on it...")

        try:
            result = agent.invoke({"messages": [("user", user_message)]})
            assistant_message = result["messages"][-1].content
        except Exception as error:
            assistant_message = f"The agent ran into an error: {error}"

        print("\n" + "-" * 54)
        print(f"AI: {assistant_message}")
        print("-" * 54)
